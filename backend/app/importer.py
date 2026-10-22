"""Excel import for the IE folder sheets. Two layouts are understood:

1. List layout (Unit 1 "Folders Detail.xlsx") - blocks of:
       <Category title>            (single text cell)
       SR# | Description | Qty     (header row)
       1   | Waist band ... | 15   (item rows; SR# may be blank)
   Columns are located from the "SR#" header, so the block can sit anywhere on the sheet.

2. Matrix layout (Unit 2 "Folder Guage set unit 2.xlsx") - blocks of:
       <Category title>
       Folders | conversion kit size | (kit qty) | Single Part | 2 part     (optional header row)
       3.2CM   | 3.2cm               | 20        | 10          | 10
   One row holds several items (folder per variant, conversion kit, gauge set...). "nil" / blank = none in stock.
"""
import re
from datetime import date

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Category, Folder, Inventory, Transaction, Unit

KNOWN_CATEGORIES = {
    # Unit 1
    "waist band": dict(name="Waist Band", code="WB", color="#2563eb",
                       description="Waist band folders: single part, two part straight, contour and piping."),
    "dnls": dict(name="DNLS", code="DN", color="#d97706",
                 description="DNLS hang folders."),
    "chain stitch 380": dict(name="Chain Stitch 380", code="CS", color="#0d9488",
                             description="Hemming folders used on chain stitch 380 machines."),
    "waist band press folder": dict(name="Waist Band Press", code="WP", color="#be123c",
                                    description="Waist band press folders and large press folders."),
    # Unit 2
    "w/b kansai": dict(name="W/B Kansai", code="KS", color="#4f46e5",
                       description="Kansai waist band folders (single part and 2 part), conversion kits and special gauge sets."),
    "heming folder": dict(name="Hemming Folder", code="HM", color="#0891b2",
                          description="Hemming folders by width."),
    "hemming folder": dict(name="Hemming Folder", code="HM", color="#0891b2",
                           description="Hemming folders by width."),
    "loop make folder+conversion kit": dict(name="Loop Make Folder", code="LM", color="#65a30d",
                                            description="Loop make folders and their conversion kits."),
    "feed of arm": dict(name="Feed Of Arm", code="FA", color="#ca8a04",
                        description="Feed-of-arm gauge sets and folders."),
    "double needle chain stitch (380)": dict(name="Double Needle Chain Stitch (380)", code="DC", color="#0d9488",
                                             description="Gauge sets for double needle chain stitch 380 machines."),
    "double needle lock stitch": dict(name="Double Needle Lock Stitch", code="DL", color="#db2777",
                                      description="Gauge sets for double needle lock stitch machines."),
    "welt-bone ass+ juki": dict(name="Welt-Bone (ASS + JUKI)", code="WJ", color="#7c3aed",
                                description="Welt-bone gauge sets, ASS and JUKI types."),
    "make loop (kansai)": dict(name="Make Loop (Kansai)", code="MK", color="#dc2626",
                               description="Make-loop gauge sets for Kansai machines."),
    "blind loop": dict(name="Blind Loop", code="BL", color="#475569",
                       description="Blind loop gauge sets."),
}
EXTRA_COLORS = ["#db2777", "#0891b2", "#65a30d", "#ca8a04", "#4f46e5", "#dc2626"]


def norm(text) -> str:
    return " ".join(str(text or "").split()).lower()


def tidy_name(text: str) -> str:
    words = " ".join(str(text).split()).split(" ")
    return " ".join(w[0].upper() + w[1:] if w.isalpha() and not w.isupper() else w for w in words)


def parse_description(desc: str) -> dict:
    low = norm(desc)
    if "single part" in low:
        ftype = "Single Part"
    elif "piping" in low:
        ftype = "Two Part Piping"
    elif "contour" in low:
        ftype = "Two Part Contour"
    elif "straight" in low:
        ftype = "Two Part Straight"
    elif "hang" in low:
        ftype = "Hang Folder"
    elif "hemming" in low:
        ftype = "Hemming Folder"
    elif "press" in low:
        ftype = "Press Folder"
    elif "gauge" in low:
        ftype = "Gauge Set"
    elif "conversion kit" in low:
        ftype = "Conversion Kit"
    else:
        ftype = "Folder"

    size, variant = "", ""
    if m := re.search(r"(\d+)\s*\*\s*(\d+/\d+)\s*([ML])?(?![A-Za-z])", desc, re.I):
        size = f'{m.group(1)}-{m.group(2)}"'
        variant = (m.group(3) or "").upper()
    elif m := re.search(r"(\d+)\s*mm", desc, re.I):
        size = f"{m.group(1)} mm"
    elif m := re.search(r"\b(\d+/\d+)\b", desc):
        size = f'{m.group(1)}"'
    elif "extra large" in low:
        size = "Extra Large"
    elif "large" in low:
        size = "Large"
    return {"type": ftype, "size": size, "variant": variant}


def _load(source):
    return openpyxl.load_workbook(source, data_only=True)


def parse_workbook(source) -> list[dict]:
    """Return [{"category": str, "items": [{"sr", "description", "qty", ...}], "skipped": int}].

    Tries the SR#/Description/Qty list layout first, then the matrix layout."""
    wb = _load(source)
    return _parse_list_layout(wb) or _parse_matrix_layout(wb)


def _parse_list_layout(wb) -> list[dict]:
    blocks: list[dict] = []
    for ws in wb.worksheets:
        pending_title = None
        cols = None  # (sr_col, desc_col, qty_col)
        current = None
        for row in ws.iter_rows(values_only=True):
            cells = [(i, v) for i, v in enumerate(row) if v is not None and str(v).strip() != ""]
            if not cells:
                continue
            sr_col = next((i for i, v in cells if norm(v).replace(" ", "") in ("sr#", "sr", "sr.", "srno", "sr.no")), None)
            if sr_col is not None:
                cols = (sr_col, sr_col + 1, sr_col + 2)
                current = {"category": (pending_title or "Uncategorized").strip(), "items": [], "skipped": 0}
                blocks.append(current)
                pending_title = None
                continue
            if len(cells) == 1 and isinstance(cells[0][1], str):
                pending_title = cells[0][1]
                current = None
                continue
            if current is not None and cols:
                sr_c, d_c, q_c = cols
                desc = row[d_c] if d_c < len(row) else None
                qty = row[q_c] if q_c < len(row) else None
                if isinstance(desc, str) and desc.strip():
                    try:
                        qty = int(float(qty or 0))
                    except (TypeError, ValueError):
                        qty = 0
                    sr = row[sr_c] if isinstance(row[sr_c], (int, float)) else None
                    current["items"].append({"sr": int(sr) if sr is not None else None,
                                             "description": " ".join(desc.split()), "qty": qty})
    return [b for b in blocks if b["items"]]


# ---------------------------------------------------------------- matrix layout (Unit 2)

_GAUGE = re.compile(r"\bg(?:au|ua)ge", re.I)
_HEADER = re.compile(r"^(qty|folders?|single part|2 part|two part|conversion kit( size)?|g(?:au|ua)ge sets?)$")


def _text(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def _qty(v) -> int | None:
    """Number of pieces in a cell: 12 -> 12, '4+4' -> 8, 'nil' -> 0, blank -> None."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    t = norm(v)
    if t in ("nil", "-", "--", "n/a", "na", "none"):
        return 0
    if re.fullmatch(r"\d+(\s*\+\s*\d+)*", t):
        return sum(int(p) for p in t.split("+"))
    return None


def _clean(text: str) -> str:
    """Fix the sheet's spelling and spacing: 'Guage' -> 'Gauge', '3.2CM' -> '3.2 cm', '1/4MH:' -> '1/4 MH'."""
    t = _text(text).strip(" :")
    t = re.sub(r"g(?:au|ua)ge", "Gauge", t, flags=re.I)
    t = re.sub(r"narr+ow", "narrow", t, flags=re.I)
    t = re.sub(r"\bne+d+le\b", "Needle", t, flags=re.I)
    t = re.sub(r"(\d)\s*(cm|mm)\b", lambda m: f"{m.group(1)} {m.group(2).lower()}", t, flags=re.I)
    t = re.sub(r"(\d)(mm|cm)(?=\d)", r"\1\2 ", t)  # 16mm5/8 -> 16 mm 5/8
    t = re.sub(r"(\d)\s*mm\s*(?=\d)", r"\1 mm ", t)
    t = re.sub(r"(\d/\d+)\s*(MH)\b", r"\1 \2", t, flags=re.I)
    t = re.sub(r"\(\s*", "(", t)
    t = re.sub(r"\s*\)", ")", t)
    t = re.sub(r"\s*\+\s*", " + ", t)
    words = []
    for w in t.split(" "):
        if w.isupper() and w.isalpha() and len(w) > 4:
            w = w.capitalize()  # NEEDLE -> Needle; keeps ASS, JUKI, MH
        elif w.isalpha() and w.islower() and w not in ("cm", "mm"):
            w = w.capitalize()
        words.append(w)
    return " ".join(words)


def _strip_kind(text: str) -> str:
    """'Special Gauge Set Narrow 6 Needle' -> 'Special Narrow 6 Needle'."""
    return " ".join(re.sub(r"\bGauge\s+Sets?\b", "", text, flags=re.I).split())


def _roles(header: list | None, width: int) -> list:
    """Work out what each column holds: ('label', kind) | ('qty', label_col, variant) | ('note',)."""
    if header is None:  # no header row: label | kit label | kit qty | label qty
        roles = [("label", "folder"), ("label", "kit"), ("qty", 1, None), ("qty", 0, None)]
        return (roles + [("note",)] * width)[:max(width, 4)]
    width = max(width, len(header))
    h = [norm(header[j]) if j < len(header) else "" for j in range(width)]
    roles: list = [None] * width
    has_qty: set[int] = set()
    for j, n in enumerate(h):
        if j == 0:
            roles[0] = ("label", "gauge" if _GAUGE.search(n) else "folder")
        elif n in ("folder", "folders"):
            roles[j] = ("label", "folder")
        elif n.startswith("conversion kit"):
            roles[j] = ("label", "kit")
        elif _GAUGE.search(n):
            roles[j] = ("label", "gauge")
        elif n == "single part":
            roles[j] = ("qty", 0, "Single Part")
            has_qty.add(0)
        elif n in ("2 part", "two part"):
            roles[j] = ("qty", 0, "Two Part")
            has_qty.add(0)
    for j, n in enumerate(h):
        if n == "qty":
            target = next((k for k in range(j - 1, -1, -1) if roles[k] and roles[k][0] == "label" and k not in has_qty), 0)
            roles[j] = ("qty", target, None)
            has_qty.add(target)
    for j in range(1, width):
        if roles[j] is None:
            left = roles[j - 1]
            if left and left[0] == "label" and (j - 1) not in has_qty:
                roles[j] = ("qty", j - 1, None)
                has_qty.add(j - 1)
            else:
                roles[j] = ("note",)
    return roles


def _block_display(title: str) -> str:
    meta = KNOWN_CATEGORIES.get(norm(title))
    return meta["name"] if meta else _clean(title)


def _matrix_items(block_title: str, roles: list, row: tuple) -> tuple[list[dict], int]:
    """Turn one sheet row into items. Returns (items, skipped_count)."""
    vals = list(row) + [None] * (len(roles) - len(row))
    block = _block_display(block_title)
    short = re.sub(r"\s*\bFolders?\b\s*$", "", block).strip() or block  # "Loop Make Folder" -> "Loop Make"
    first = _text(vals[0])
    special = roles[0][1] == "folder" and bool(_GAUGE.search(first))  # e.g. "Special Guage Set narrow 6 needle"
    notes = [_text(vals[j]) for j, r in enumerate(roles) if r[0] == "note" and isinstance(vals[j], str) and vals[j].strip()]
    note = " ".join(notes)
    variant_note = _clean(re.sub(r"\btype\b", "", note, flags=re.I)).upper() if note else ""
    items, skipped = [], 0

    for j, role in enumerate(roles):
        if role[0] != "qty":
            continue
        raw = vals[j]
        qty = _qty(raw)
        if qty is None and raw is None:
            continue
        label_col, variant = role[1], role[2]
        kind = roles[label_col][1] if roles[label_col][0] == "label" else "folder"
        label = _text(vals[label_col])
        if not label and label_col != 0:
            label = first  # e.g. multi-needle row: kit qty with no kit size -> describes the col-0 item
            kind = "gauge" if _GAUGE.search(first) else kind
        if not label or norm(label) in ("nil", "na"):
            skipped += 1
            continue
        if not qty:
            skipped += 1
            continue

        if special:
            base = _clean(first)
            size = _clean(vals[1]) if len(vals) > 1 and _text(vals[1]) else ""
            size = size if not _GAUGE.search(size) else ""
            if kind == "folder" or variant:
                name = " ".join(filter(None, [_strip_kind(base), variant, "Folder", size]))
                ftype = variant or "Folder"
            else:
                name = " ".join(filter(None, [base, size]))
                ftype = "Gauge Set"
            size_out = size
        else:
            size_out = _clean(re.sub(r"^\s*conversion kit\s*", "", label, flags=re.I))
            if kind == "kit":
                name, ftype = f"{short} Conversion Kit {size_out}", "Conversion Kit"
            elif kind == "gauge":
                name, ftype = f"{short} Gauge Set {size_out}", "Gauge Set"
            else:
                if variant:
                    name, ftype = f"{short} {variant} Folder {size_out}", variant
                else:
                    name = f"{block} {size_out}" if re.search(r"\bfolder\b", block, re.I) else f"{block} Folder {size_out}"
                    ftype = "Hemming Folder" if "hem" in norm(block) else "Loop Folder" if "loop" in norm(block) else "Folder"
        name = " ".join(name.split())
        item_variant = variant_note if (kind == "gauge" and variant_note) else ""
        if item_variant:
            name = f"{name} ({item_variant})"
        specs = f"Quantity in sheet: {_text(raw)}" if isinstance(raw, str) and "+" in raw else ""
        items.append({"sr": None, "description": name, "qty": qty, "name": name, "type": ftype,
                      "size": size_out, "variant": item_variant, "specs": specs})
    return items, skipped


def _parse_matrix_layout(wb) -> list[dict]:
    blocks: list[dict] = []
    for ws in wb.worksheets:
        current = None
        roles = None
        for row in ws.iter_rows(values_only=True):
            cells = [(i, v) for i, v in enumerate(row) if v is not None and str(v).strip() != ""]
            if not cells:
                continue
            if len(cells) == 1 and isinstance(cells[0][1], str) and _qty(cells[0][1]) is None:
                current = {"category": _text(cells[0][1]), "items": [], "skipped": 0}
                blocks.append(current)
                roles = _roles(None, len(row))
                continue
            if current is None:
                continue
            if all(isinstance(v, str) for _, v in cells) and all(_HEADER.match(norm(v)) for _, v in cells):
                roles = _roles(list(row), len(row))
                continue
            if row[0] is None:
                continue
            items, skipped = _matrix_items(current["category"], roles, row)
            current["skipped"] += skipped
            seen = {norm(i["description"]) for i in current["items"]}
            for it in items:
                if norm(it["description"]) in seen:  # same item listed twice: add up
                    next(i for i in current["items"] if norm(i["description"]) == norm(it["description"]))["qty"] += it["qty"]
                else:
                    current["items"].append(it)
                    seen.add(norm(it["description"]))
    return [b for b in blocks if b["items"]]


# ---------------------------------------------------------------- database upsert

def _get_or_create_category(db: Session, unit: Unit | None, title: str, order: int) -> Category:
    key = norm(title)
    meta = KNOWN_CATEGORIES.get(key)
    unit_id = unit.id if unit else None
    for cat in db.scalars(select(Category).where(Category.unit_id == unit_id)):
        if norm(cat.source_name) == key or norm(cat.name) == key or (meta and cat.name == meta["name"]):
            return cat
    existing = {c.code for c in db.scalars(select(Category))}
    used = len(existing)
    if meta:
        meta = dict(meta)
    else:
        meta = dict(name=_clean(title) if unit and unit.item_prefix else tidy_name(title), code="",
                    color=EXTRA_COLORS[used % len(EXTRA_COLORS)], description="")
    if not meta["code"] or meta["code"] in existing:
        base = meta["code"] or "".join(w[0] for w in re.findall(r"[A-Za-z]+", meta["name"])[:3]).upper() or "XX"
        code, n = base, 2
        while code in existing:
            code, n = f"{base}{n}", n + 1
        meta["code"] = code
    cat = Category(unit_id=unit_id, source_name=title.strip(), sort_order=order, **meta)
    db.add(cat)
    db.flush()
    return cat


def item_prefix(cat: Category) -> str:
    return f"{cat.unit.item_prefix if cat.unit else ''}{cat.code}-"


def next_item_code(db: Session, cat: Category) -> str:
    prefix = item_prefix(cat)
    codes = db.scalars(select(Folder.item_code).where(Folder.item_code.like(f"{prefix}%"))).all()
    nums = [int(c[len(prefix):]) for c in codes if c[len(prefix):].isdigit()]
    return f"{prefix}{(max(nums) + 1 if nums else 1):03d}"


def import_blocks(db: Session, blocks: list[dict], user_id: int | None = None, unit: Unit | None = None) -> dict:
    """Upsert categories/folders into a unit. Existing folders get their total synced to the sheet quantity."""
    stats = {"categories": 0, "created": 0, "updated": 0, "unchanged": 0}
    today = date.today()
    source = f" ({unit.name})" if unit else ""
    for order, block in enumerate(blocks):
        cat = _get_or_create_category(db, unit, block["category"], order)
        stats["categories"] += 1
        existing = {norm(f.description): f for f in cat.folders}
        for item in block["items"]:
            folder = existing.get(norm(item["description"]))
            if folder is None:
                parsed = parse_description(item["description"])
                if "type" in item:  # matrix layout already knows type/size/variant
                    parsed = {k: item[k] for k in ("type", "size", "variant")}
                folder = Folder(category=cat, name=item.get("name") or tidy_name(item["description"]),
                                description=item["description"], item_code=next_item_code(db, cat), sr_no=item["sr"],
                                specs=item.get("specs", ""), **parsed)
                folder.inventory = Inventory(total_qty=item["qty"], available_qty=item["qty"])
                db.add(folder)
                db.flush()
                db.add(Transaction(folder_id=folder.id, type="import", quantity=item["qty"], user_id=user_id,
                                   date=today, remarks=f"Imported from Excel{source}"))
                existing[norm(item["description"])] = folder
                stats["created"] += 1
                continue
            inv = folder.inventory
            delta = item["qty"] - inv.total_qty
            if delta == 0:
                stats["unchanged"] += 1
                continue
            inv.total_qty = item["qty"]
            inv.available_qty = max(0, inv.available_qty + delta)
            inv.total_qty = max(inv.total_qty, inv.available_qty + inv.issued_qty + inv.damaged_qty)
            db.add(Transaction(folder_id=folder.id, type="adjust", quantity=delta, user_id=user_id, date=today,
                               remarks=f"Excel re-import{source}: total set to {item['qty']}"))
            stats["updated"] += 1
    db.commit()
    return stats
