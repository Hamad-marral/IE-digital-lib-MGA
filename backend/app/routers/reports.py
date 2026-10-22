import io
from datetime import date, datetime

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select, true
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Category, Folder, Transaction, Unit
from ..security import current_user
from ..serializers import folder_out, stock_status

router = APIRouter(prefix="/api/reports", tags=["reports"], dependencies=[Depends(current_user)])

STATUS_LABEL = {"available": "Available", "low_stock": "Low Stock", "out_of_stock": "Out of Stock"}


def _unit_where(unit_id: int | None):
    return (Category.unit_id == unit_id) if unit_id else true()


def _folders(db: Session, unit_id: int | None):
    return db.scalars(select(Folder).join(Category).outerjoin(Unit).where(_unit_where(unit_id))
                      .order_by(Unit.sort_order, Category.sort_order, Folder.item_code)).all()


def _unit_name(f: Folder) -> str:
    return f.category.unit.name if f.category.unit else ""


@router.get("/summary")
def summary(date_from: date | None = None, date_to: date | None = None, unit_id: int | None = None,
            db: Session = Depends(get_db)):
    tq = select(Transaction).join(Folder, Transaction.folder_id == Folder.id).join(Category).where(_unit_where(unit_id))
    if date_from:
        tq = tq.where(Transaction.date >= date_from)
    if date_to:
        tq = tq.where(Transaction.date <= date_to)
    txs = db.scalars(tq).all()

    by_line: dict[str, dict] = {}
    by_folder: dict[int, dict] = {}
    by_month: dict[str, dict] = {}
    for t in txs:
        if t.type not in ("issue", "return"):
            continue
        line = by_line.setdefault(t.to_from or "—", {"line": t.to_from or "—", "issued": 0, "returned": 0})
        line["issued" if t.type == "issue" else "returned"] += t.quantity
        month = by_month.setdefault(t.date.strftime("%Y-%m"), {"month": t.date.strftime("%b %Y"), "key": t.date.strftime("%Y-%m"), "issued": 0, "returned": 0})
        month["issued" if t.type == "issue" else "returned"] += t.quantity
        if t.type == "issue":
            f = by_folder.setdefault(t.folder_id, {"folder_id": t.folder_id, "name": t.folder.name,
                                                    "item_code": t.folder.item_code, "issued": 0, "times": 0})
            f["issued"] += t.quantity
            f["times"] += 1

    categories = []
    for c in db.scalars(select(Category).outerjoin(Unit).where(_unit_where(unit_id))
                        .order_by(Unit.sort_order, Category.sort_order)):
        inv = [f.inventory for f in c.folders]
        categories.append({"id": c.id, "name": c.name, "code": c.code, "unit": c.unit.name if c.unit else "", "color": c.color, "types": len(c.folders),
                           "total": sum(i.total_qty for i in inv), "available": sum(i.available_qty for i in inv),
                           "issued": sum(i.issued_qty for i in inv), "damaged": sum(i.damaged_qty for i in inv)})

    alerts = [folder_out(f) for f in _folders(db, unit_id) if stock_status(f.inventory) != "available"]
    alerts.sort(key=lambda f: (f["available_qty"], f["name"]))
    return {
        "categories": categories,
        "by_line": sorted(by_line.values(), key=lambda r: -r["issued"]),
        "top_folders": sorted(by_folder.values(), key=lambda r: -r["issued"])[:10],
        "by_month": sorted(by_month.values(), key=lambda r: r["key"]),
        "alerts": alerts,
        "transaction_count": len(txs),
    }


def _sheet(ws, headers: list[str], rows: list[list]):
    head_fill = PatternFill("solid", fgColor="1E3A6E")
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] or [10])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _file_name(db: Session, base: str, unit_id: int | None) -> str:
    unit = db.get(Unit, unit_id) if unit_id else None
    return f"{base}_{unit.name.replace(' ', '')}" if unit else f"{base}_AllUnits"


def _xlsx(wb: Workbook, name: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}_{stamp}.xlsx"'})


@router.get("/export/inventory")
def export_inventory(unit_id: int | None = None, db: Session = Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    folders = _folders(db, unit_id)
    rows = [[f.item_code, _unit_name(f), f.category.name, f.name, f.type, f.size, f.variant, f.inventory.total_qty,
             f.inventory.available_qty, f.inventory.issued_qty, f.inventory.damaged_qty, f.inventory.min_qty,
             STATUS_LABEL[stock_status(f.inventory)], f.inventory.location] for f in folders]
    _sheet(ws, ["Item Code", "Unit", "Category", "Folder", "Type", "Size", "Variant", "Total", "Available", "Issued",
                "Damaged", "Min Stock", "Status", "Location"], rows)
    return _xlsx(wb, _file_name(db, "IE_Inventory", unit_id))


@router.get("/export/transactions")
def export_transactions(date_from: date | None = None, date_to: date | None = None, unit_id: int | None = None,
                        db: Session = Depends(get_db)):
    q = (select(Transaction).join(Folder, Transaction.folder_id == Folder.id).join(Category)
         .where(_unit_where(unit_id)).order_by(Transaction.date.desc(), Transaction.id.desc()))
    if date_from:
        q = q.where(Transaction.date >= date_from)
    if date_to:
        q = q.where(Transaction.date <= date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    rows = [[t.date.isoformat(), t.type.replace("_", " ").title(), _unit_name(t.folder), t.folder.item_code, t.folder.name, t.quantity,
             t.to_from, t.line_machine, t.receiver, t.condition, t.user.name if t.user else "", t.remarks]
            for t in db.scalars(q)]
    _sheet(ws, ["Date", "Type", "Unit", "Item Code", "Folder", "Qty", "Line", "Machine", "Receiver", "Condition",
                "Recorded By", "Remarks"], rows)
    return _xlsx(wb, _file_name(db, "IE_Transactions", unit_id))


@router.get("/export/low-stock")
def export_low_stock(unit_id: int | None = None, db: Session = Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Low Stock"
    rows = [[f.item_code, _unit_name(f), f.category.name, f.name, f.inventory.available_qty, f.inventory.min_qty,
             f.inventory.total_qty, STATUS_LABEL[stock_status(f.inventory)]]
            for f in _folders(db, unit_id) if stock_status(f.inventory) != "available"]
    rows.sort(key=lambda r: r[4])
    _sheet(ws, ["Item Code", "Unit", "Category", "Folder", "Available", "Min Stock", "Total", "Status"], rows)
    return _xlsx(wb, _file_name(db, "IE_Low_Stock", unit_id))
